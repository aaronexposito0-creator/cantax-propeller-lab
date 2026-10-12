"""Meaningful local checks; saved evidence tests never claim a new native run."""
import copy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from engineering_checks import coefficients, ideal_disk_power, preflight_xlsx, read, validate_case, verify_saved_case
from build_workbook import build
from mesh_and_structure import analyze_mesh, gmsh_path_literal
from run_lab import family, refresh_panel
from runtime_config import executable, local_engine_environment
from openpyxl import load_workbook

class ScientificChecks(unittest.TestCase):
    def test_saved_examples_independent_definitions(self):
        for name in ('baseline','rpm_1000','five_blades','diameter_study'):
            with self.subTest(example=name):
                proof=verify_saved_case(ROOT/'examples/verified_runs'/name)
                self.assertEqual(proof['status'],'PASS')
                self.assertGreaterEqual(proof['numeric_identity_checks'],60)

    def test_power_and_efficiency_dimensions(self):
        values=coefficients(900,2,.30,1025,350.348984,13.555366)
        self.assertAlmostEqual(values['J'],4/9)
        self.assertAlmostEqual(values['P_W'],1277.56313,places=3)
        self.assertAlmostEqual(values['eta0'],values['J']*values['KT']/(2*__import__('math').pi*values['KQ']))
        self.assertGreater(values['P_W'],ideal_disk_power(350.348984,2,.30,1025))

    def test_dimensionally_similar_operating_points(self):
        a=coefficients(900,2,.3,1025,350,13)
        b=coefficients(450,2,.6,1025,1400,104)
        for key in ('J','KT','KQ','eta0'):
            self.assertAlmostEqual(a[key],b[key])

    def test_guard_invalid_inputs(self):
        for value in (0,-100,float('nan'),True):
            with self.subTest(rpm=value),self.assertRaises(ValueError):
                coefficients(value,2,.3,1025,350,13)
        original=read(ROOT/'config/default_case.json')
        for mutation in ('bad_blades','bad_material','repeated_radial','nonfinite'):
            case=copy.deepcopy(original)
            if mutation=='bad_blades':case['geometry']['blades']=3.5
            elif mutation=='bad_material':case['material']='UNKNOWN'
            elif mutation=='repeated_radial':case['radial_distribution'][1]['r_R']=case['radial_distribution'][0]['r_R']
            else:case['geometry']['diameter_m']=float('inf')
            with self.subTest(mutation=mutation),self.assertRaises(ValueError):validate_case(case)

    def test_sweep_is_bounded_and_distinct(self):
        cases=family(read(ROOT/'config/default_case.json'),diameter_sweep=True)
        self.assertEqual(len(cases),13)
        self.assertEqual(len({entry['id'] for entry in cases}),13)
        self.assertEqual({entry['geometry']['blades'] for entry in cases},{3,4,5})
        self.assertEqual(len({entry['geometry']['diameter_m'] for entry in cases}),3)

class WorkbookAndOutputChecks(unittest.TestCase):
    def setUp(self):
        (ROOT/'runs').mkdir(exist_ok=True)
        self.temporary=tempfile.TemporaryDirectory(dir=ROOT/'runs')
        self.directory=Path(self.temporary.name)
    def tearDown(self):self.temporary.cleanup()

    def test_native_formula_schema_and_no_active_content(self):
        path=build(self.directory/'input.xlsx')
        self.assertTrue(preflight_xlsx(path))
        book=load_workbook(path,data_only=False)
        self.assertEqual(book.sheetnames,['Panel','Control','Radial','Alternativas','Curvas','Materiales','Verificacion'])
        self.assertEqual(book['Control']['B14'].value,900)
        self.assertEqual(len(book['Control'].data_validations.dataValidation),25)
        self.assertEqual(book['Radial']['J6'].value,'=SUM(H6:H13)')
        self.assertIn('Control!$B$7',book['Radial']['G6'].value)
        book.close()
        with self.assertRaises(ValueError):build(path)

    def test_external_and_dde_formula_rejected(self):
        for index,formula in enumerate(('WEBSERVICE("https://example.invalid")',"cmd|' /C something'!A1","'[1]Sheet'!A1",'RTD("server",,"secret")')):
            path=self.directory/f'bad_{index}.xlsx'
            with ZipFile(path,'w') as archive:
                archive.writestr('xl/worksheets/sheet1.xml',f'<worksheet><f>{formula}</f></worksheet>')
            with self.subTest(formula=formula),self.assertRaises(ValueError):preflight_xlsx(path)

    def test_changed_controller_and_locked_target_preserve_results(self):
        target=self.directory/'controller.xlsx'
        target.write_bytes(b'saved inputs')
        digest=hashlib.sha256(target.read_bytes()).hexdigest()
        (self.directory/'CantaX_Propeller_Resultados.xlsx').write_bytes(b'computed results')
        target.write_bytes(b'new user edits')
        result=refresh_panel(self.directory,target,digest,target)
        self.assertEqual(result['status'],'DEFERRED_CONTROLLER_CHANGED')
        self.assertEqual(target.read_bytes(),b'new user edits')
        self.assertEqual((self.directory/'panel_refresh.xlsx').read_bytes(),b'computed results')
        digest=hashlib.sha256(target.read_bytes()).hexdigest()
        with patch('run_lab.os.replace',side_effect=PermissionError('controlled lock')):
            result=refresh_panel(self.directory,target,digest,target)
        self.assertEqual(result['status'],'DEFERRED_PANEL_LOCKED')
        self.assertEqual(target.read_bytes(),b'new user edits')

    def test_tetrahedron_orientation_and_volume_acceptance(self):
        content='$MeshFormat\n2.2 0 8\n$EndMeshFormat\n$Nodes\n4\n1 0 0 0\n2 1 0 0\n3 0 1 0\n4 0 0 1\n$EndNodes\n$Elements\n1\n1 4 0 {nodes}\n$EndElements\n'
        path=self.directory/'one.msh'
        path.write_text(content.format(nodes='1 2 3 4'),encoding='ascii')
        result,*_=analyze_mesh(path,1/6)
        self.assertEqual(result['status'],'PASS')
        self.assertEqual(result['boundary_triangles'],4)
        path.write_text(content.format(nodes='1 3 2 4'),encoding='ascii')
        result,*_=analyze_mesh(path,1/6)
        self.assertEqual(result['status'],'FAIL')
        self.assertEqual(result['negative_orientation_elements'],1)

    def test_gmsh_path_cannot_inject_geo_commands(self):
        self.assertEqual(gmsh_path_literal(Path('cad/BASE.step')),'"cad/BASE.step"')
        for path in ('model"; System("command"); //','model\nMerge "other.step";','model\r.step'):
            with self.subTest(path=path),self.assertRaises(ValueError):gmsh_path_literal(Path(path))

    def test_explicit_missing_executable_does_not_fall_back(self):
        with patch.dict('os.environ',{'CANTAX_GMSH':str(self.directory/'missing.exe')}):
            with self.assertRaises(FileNotFoundError):executable('gmsh')

    def test_engine_environment_excludes_service_tokens(self):
        values={'OPENAI_API_KEY':'synthetic','ANTHROPIC_API_KEY':'synthetic','GH_TOKEN':'synthetic','GITHUB_TOKEN':'synthetic','CODEX_HOME':'synthetic','CHATGPT_TOKEN':'synthetic','LM_LICENSE_FILE':'27000@local-license-host','MLM_LICENSE_FILE':'local.lic','CANTAX_EXECUTABLES':'config/executables.local.json','PATH':'local-bin'}
        filtered=local_engine_environment(values)
        self.assertEqual(set(filtered),{'LM_LICENSE_FILE','MLM_LICENSE_FILE','CANTAX_EXECUTABLES','PATH'})

if __name__=='__main__':unittest.main()
