"""Original macro-free workbook authoring using the public openpyxl API.

This builder does not evaluate formulas. Microsoft Excel COM evaluates and
independently checks them during a full run; the saved-example demo says so.
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from engineering_checks import read, validate_case

ROOT = Path(__file__).resolve().parents[1]
NAVY = "15304B"
AMBER = "FFF0C7"

def build(output, case=None, run=None):
    output = Path(output)
    if not output.resolve().is_relative_to(ROOT) or output.exists():
        raise ValueError("Workbook output must be a fresh file within this checkout")
    case = validate_case(case or read(ROOT / "config/default_case.json"))
    material_file = Path(run)/"config_snapshot/materials.json" if run else ROOT/"config/materials.json"
    if not material_file.is_file():
        material_file = ROOT/"config/materials.json"
    materials = read(material_file)["materials"]
    mapping = read(ROOT/"config/excel_mapping.json")["parameters"]
    combined = read(Path(run)/"combined_results.json") if run else {}
    hydro = read(Path(run)/"hydro_results.json") if run else {}
    dynamics = read(Path(run)/"dynamics_result.json") if run else {}
    book = Workbook()
    book.remove(book.active)
    sheets = {name: book.create_sheet(name) for name in ("Panel", "Control", "Radial", "Alternativas", "Curvas", "Materiales", "Verificacion")}
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    book.properties.creator = "CantaX Propeller Lab contributors"
    book.properties.title = "Parametric marine propeller preliminary engineering"
    for sheet in sheets.values():
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B6"
        sheet.column_dimensions["A"].width = 44
        for col in "BCDEFGHIJKLMNOPQ":
            sheet.column_dimensions[col].width = 21
        sheet["A1"] = "CantaX Propeller Lab — " + sheet.title
        sheet["A1"].font = Font(name="Arial", size=18, bold=True, color=NAVY)
        sheet.row_dimensions[1].height = 31

    def row(sheet, index, values, header=False):
        for col, value in enumerate(values, 1):
            cell = sheet.cell(index, col, value)
            cell.font = Font(name="Arial", size=10, color="FFFFFF" if header else "273B4C", bold=header)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if header:
                cell.fill = PatternFill("solid", fgColor=NAVY)
        sheet.row_dimensions[index].height = 32 if header else 28

    control = sheets["Control"]
    control["A2"] = "Amber inputs. Save and close Excel before launching; all inputs are explicit scenarios."
    row(control, 5, ["Parameter", "Value", "Unit", "Minimum", "Maximum", "Input check", "Use", "Run input", "Changed since run"], True)
    for index, parameter in enumerate(mapping, 6):
        group, key = parameter["group"], parameter["key"]
        value = case[group][key] if group else case[key]
        row(control, index, [key, value, parameter["unit"], parameter["min"], parameter["max"], None, "SCENARIO", value if run else None])
        control.cell(index, 2).fill = PatternFill("solid", fgColor=AMBER)
        if parameter["min"] is not None:
            validation = DataValidation(type="whole" if key=="blades" else "decimal", operator="between", formula1=parameter["min"], formula2=parameter["max"])
            control.cell(index, 6, f'=IF(B{index}="","MISSING",IF(OR(B{index}<D{index},B{index}>E{index}),"OUT OF RANGE","OK"))')
        else:
            validation = DataValidation(type="list", formula1='"'+','.join(m["id"] for m in materials)+'"')
            control.cell(index, 6, f'=IF(B{index}="","MISSING","CHECK MATERIAL")')
        validation.allow_blank = False
        validation.showErrorMessage = True
        validation.error = "Choose a value within the demonstration bounds."
        control.add_data_validation(validation)
        validation.add(control.cell(index, 2))
        if run:
            control.cell(index, 9, f'=IF(B{index}=H{index},"OK","MODIFICADO")')
        control.cell(index, 2).number_format = "0.00E+00" if "viscosity" in key else "0.000"

    radial = sheets["Radial"]
    radial["A2"] = "Nominal strip chord normalization; CAD tip clipping is accounted for separately."
    row(radial, 5, ["r/R", "Chord weight", "Pitch factor", "Thickness factor", "Skew factor", "Rake factor", "Chord (m)", "Integral (m)"], True)
    keys = ("r_R", "chord_weight", "pitch_factor", "thickness_factor", "skew_factor", "rake_factor")
    for index, station in enumerate(case["radial_distribution"], 6):
        for col, key in enumerate(keys, 1):
            radial.cell(index, col, station[key]).fill = PatternFill("solid", fgColor=AMBER)
            if run:
                radial.cell(index, col+10, station[key])
        radial.cell(index, 7, f'=B{index}*Control!$B$9*PI()*Control!$B$6^2/(4*Control!$B$7*$J$6)')
        radial.cell(index, 8, f'=(A{index+1}-A{index})*(B{index+1}+B{index})/2*Control!$B$6/2' if index<13 else '=0')
        if run:
            conditions = ','.join(f'{a}{index}={b}{index}' for a, b in zip("ABCDEF", "KLMNOP"))
            radial.cell(index, 17, f'=IF(AND({conditions}),"OK","MODIFICADO")')
        for cell in radial[index]:
            cell.number_format = "0.0000"
    radial["J5"] = "Normalization integral"
    radial["J6"] = "=SUM(H6:H13)"

    panel = sheets["Panel"]
    panel["A2"] = "Calculated preliminary estimates. No measured polar, certified naval assessment or validated 3D CFD."
    baseline = hydro.get("baseline", {})
    metrics = [
        ["Execution", Path(run).name if run else "NOT RUN", "", "Saved run"],
        ["Diameter", case["geometry"]["diameter_m"], "m", "Scenario"],
        ["Thrust", baseline.get("T_N", "NOT CALCULATED"), "N", "Preliminary BEM"],
        ["Torque", baseline.get("Q_Nm", "NOT CALCULATED"), "N m", "Preliminary BEM"],
        ["Shaft power", baseline.get("P_W", "NOT CALCULATED"), "W", "Preliminary BEM"],
        ["Open-water efficiency", baseline.get("eta0", "NOT CALCULATED"), "1", "Preliminary BEM"],
        ["CAD mass", combined.get("baseline", {}).get("mass_kg", "NOT CALCULATED"), "kg", "Geometry and typical density"],
        ["Screening recommendation", combined.get("recommendation", {}).get("id") or "NOT RANKED", "", "Hydrodynamic/mass/CAD only"],
        ["Structural status", combined.get("structure_status", "NOT CALCULATED"), "", "Preliminary elastic estimate"],
        ["Physical validation", "UNAVAILABLE", "", "Requires experiments"],
        ["Freshness", "NOT RUN", "", "Edits require a new run"],
        ["Dynamic scenario", dynamics.get("scenario_outcome", "NOT CALCULATED"), "", "No forced tracking pass"],
        ["Final requested torque", dynamics.get("requested_steady_torque_Nm", "NOT CALCULATED"), "N m", "Check motor capacity"],
    ]
    row(panel, 5, ["Base result", "Value", "Unit", "Scope"], True)
    for index, values in enumerate(metrics, 6):
        row(panel, index, values)
    if run:
        panel["B16"] = '=IF(COUNTIFS(Control!$I$6:$I$30,"MODIFICADO")+COUNTIFS(Radial!$Q$6:$Q$13,"MODIFICADO")>0,"ENTRADAS MODIFICADAS","RESULTADOS VIGENTES")'
    panel.column_dimensions["B"].width = 44
    panel.column_dimensions["D"].width = 46
    panel["A21"] = "Edit Control/Radial → save/close Excel → Start_CantaX.bat → full run."
    panel["A22"] = "No macros, external data connections, remote calls or changes to Office security policy."

    alternatives = sheets["Alternativas"]
    row(alternatives, 5, ["ID", "Blades", "P/D", "EAR", "D (m)", "T (N)", "Q (N m)", "P (W)", "Efficiency", "Mass (kg)", "BEM domain", "CAD valid", "Screening feasible", "Rank", "Elastic analysis"], True)
    for index, entry in enumerate(combined.get("alternatives", []), 6):
        row(alternatives, index, [entry.get(key) for key in ("id", "blades", "pitch_ratio", "expanded_area_ratio", "diameter_m", "T_N", "Q_Nm", "P_W", "eta0", "mass_kg", "valid_domain", "cad_valid", "screening_feasible", "screening_rank", "structure_status")])

    curves = sheets["Curvas"]
    row(curves, 5, ["J", "KT", "KQ", "Efficiency", "Advance (m/s)", "RPM"], True)
    curve_rows = combined.get("curve_rows", [])
    for index, entry in enumerate(curve_rows, 6):
        row(curves, index, [entry.get("J"), *[entry.get(k) if entry["valid_domain"] else None for k in ("KT", "KQ", "eta0")], entry.get("Va_mps"), entry.get("rpm")])
    if curve_rows:
        chart = ScatterChart()
        chart.title = "Estimated open-water efficiency"
        chart.x_axis.title = "Advance coefficient J"
        chart.y_axis.title = "Efficiency"
        chart.series.append(Series(Reference(curves, min_col=4, min_row=6, max_row=5+len(curve_rows)), Reference(curves, min_col=1, min_row=6, max_row=5+len(curve_rows)), title="Analytical BEM"))
        chart.legend = None
        chart.height, chart.width = 12, 21
        curves.add_chart(chart, "H5")

    material_sheet = sheets["Materiales"]
    row(material_sheet, 5, ["ID", "Material", "Density (kg/m3)", "E (GPa)", "Poisson", "Yield (MPa)", "Classification", "Source"], True)
    for index, material in enumerate(materials, 6):
        row(material_sheet, index, [material["id"], material.get("label", material["id"]), material["density_kgm3"], material["elastic_modulus_Pa"]/1e9 if material.get("elastic_modulus_Pa") else "UNAVAILABLE", material.get("poisson_ratio") if material.get("poisson_ratio") is not None else "UNAVAILABLE", material["yield_strength_Pa"]/1e6 if material.get("yield_strength_Pa") else "UNAVAILABLE", "Typical public property; not certified", " ".join(material.get("source_urls", []))])
        material_sheet.row_dimensions[index].height = 60
    material_sheet.column_dimensions["H"].width = 64
    verification = sheets["Verificacion"]
    row(verification, 5, ["Stage", "Status", "Evidence / scope"], True)
    for index, values in enumerate(combined.get("checks", [["Integration", "NOT EXECUTED", "Run the configured native engines to generate new results"]]), 6):
        row(verification, index, values)
    verification.column_dimensions["C"].width = 100
    for sheet in sheets.values():
        sheet.auto_filter.ref = f"A5:{sheet.cell(5, sheet.max_column).column_letter}{max(5, sheet.max_row)}"
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    book.close()
    return output

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-dir")
    parser.add_argument("--output")
    args = parser.parse_args()
    run = Path(args.run_dir).resolve() if args.run_dir else None
    output = Path(args.output).resolve() if args.output else (run/"CantaX_Propeller_Resultados.xlsx" if run else ROOT/"CantaX_Propeller_Lab.xlsx")
    if run and not run.is_relative_to(ROOT/"runs"):
        raise ValueError("Run must remain inside this checkout's runs directory")
    result = build(output, read(run/"case.json") if run else None, run)
    print(json.dumps({"output": str(result), "authoring": "openpyxl", "formula_evaluation": "Requires native Excel; not evaluated by openpyxl"}))

if __name__ == "__main__":
    main()
